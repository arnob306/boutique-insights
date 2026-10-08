import shutil
from datetime import date

import pandas as pd
import pytest

from src.adapters.boutique_xlsx import load_workbook_data
from src.decisions.store import (
    Recommendation,
    add_recommendations,
    list_outcomes,
    list_recommendations,
    open_log,
)
from src.metrics.calendar import DEFAULT_CALENDAR_PATH, load_calendar
from src.metrics.patterns import festival_windows
from src.reminders import (
    PRIVACY_NOTE,
    build_festival_reminders,
    build_reminders,
    load_customer_rows,
)
from src.reports.delivery import DeliveryError
from src.run import main

TODAY = '2026-10-01'
SALT = 'unit-test-salt-0123456789'
SMTP_ENV = {
    'SMTP_HOST': 'smtp.example.com', 'SMTP_PORT': '587',
    'SMTP_USER': 'sender@example.com', 'SMTP_PASSWORD': 'app-password-123',
    'REPORT_TO': 'mum@example.com',
}


class Recorder:
    def __init__(self, error=None):
        self.sent = []
        self.error = error

    def __call__(self, settings, message):
        if self.error:
            raise self.error
        self.sent.append(message)


@pytest.fixture
def private_root(tmp_path):
    root = tmp_path / 'private'
    (root / 'inbox').mkdir(parents=True)
    return root


def _run(*args, sender=None):
    return main(['--today', TODAY, *args], sender=sender or Recorder(),
                env_file=None)


def _private(private_root, *extra, sender=None):
    return _run('--profile', 'private', '--private-root', str(private_root),
                *extra, sender=sender)


@pytest.fixture
def with_salt(monkeypatch):
    monkeypatch.setenv('CUSTOMER_HASH_SALT', SALT)


@pytest.fixture
def inbox_file(private_root, workbook_path):
    target = private_root / 'inbox' / 'Boutique_Sales.xlsx'
    shutil.copy(workbook_path, target)
    return target


def test_synthetic_profile_writes_a_demo_report(workbook_path, tmp_path):
    out = tmp_path / 'demo'
    code = _run('--profile', 'synthetic', '--file', str(workbook_path),
                '--demo-output', str(out))
    assert code == 0
    html = next(out.glob('weekly_*.html')).read_text(encoding='utf-8')
    assert 'Boutique weekly summary' in html
    assert next(out.glob('weekly_*.txt')).exists()


def test_private_profile_refuses_to_run_without_a_salt(
        private_root, inbox_file, monkeypatch, capsys):
    monkeypatch.delenv('CUSTOMER_HASH_SALT', raising=False)
    assert _private(private_root) == 2
    assert 'CUSTOMER_HASH_SALT' in capsys.readouterr().out
    assert not (private_root / 'output').exists()


def test_private_run_writes_output_and_archives_the_workbook(
        private_root, inbox_file, with_salt):
    assert _private(private_root) == 0
    assert list((private_root / 'output').glob('weekly_*.html'))
    assert not inbox_file.exists()
    archived = list((private_root / 'archive').glob('*.xlsx'))
    assert [p.name for p in archived] == [f'{TODAY}_Boutique_Sales.xlsx']


def test_empty_inbox_gives_a_clear_message(private_root, with_salt, capsys):
    assert _private(private_root) == 2
    assert 'inbox' in capsys.readouterr().out.lower()


def test_unreadable_workbook_is_reported_without_a_traceback(
        private_root, with_salt, capsys, tmp_path):
    from openpyxl import Workbook
    broken = private_root / 'inbox' / 'broken.xlsx'
    Workbook().save(broken)
    assert _private(private_root) == 2
    out = capsys.readouterr().out
    assert 'Traceback' not in out
    assert broken.exists()  # not archived when it failed


def test_a_file_that_is_not_a_workbook_is_reported_without_a_traceback(
        private_root, with_salt, capsys):
    broken = private_root / 'inbox' / 'broken.xlsx'
    broken.write_text('not a workbook', encoding='utf-8')
    assert _private(private_root) == 2
    out = capsys.readouterr().out
    assert 'broken.xlsx could not be opened as an Excel workbook' in out
    assert 'engine' not in out  # no library jargon
    assert broken.exists()  # not archived when it failed


def test_untrusted_data_is_never_emailed_or_archived(
        private_root, with_salt, make_workbook, monkeypatch):
    for key, value in SMTP_ENV.items():
        monkeypatch.setenv(key, value)
    bad = make_workbook(problems={'duplicate_id'})
    target = private_root / 'inbox' / 'bad.xlsx'
    shutil.copy(bad, target)
    sender = Recorder()
    assert _private(private_root, '--send', sender=sender) == 1
    assert sender.sent == []
    assert target.exists()
    assert list((private_root / 'output').glob('weekly_*.html'))


def test_send_emails_the_report(private_root, inbox_file, with_salt, monkeypatch):
    for key, value in SMTP_ENV.items():
        monkeypatch.setenv(key, value)
    sender = Recorder()
    assert _private(private_root, '--send', sender=sender) == 0
    assert len(sender.sent) == 1
    assert sender.sent[0]['To'] == 'mum@example.com'
    assert 'Boutique weekly summary' in sender.sent[0]['Subject']


def test_send_without_smtp_settings_fails_cleanly(
        private_root, inbox_file, with_salt, monkeypatch, capsys):
    for key in SMTP_ENV:
        monkeypatch.delenv(key, raising=False)
    assert _private(private_root, '--send') == 3
    assert 'SMTP_HOST' in capsys.readouterr().out


def test_delivery_failure_returns_an_error_code(
        private_root, inbox_file, with_salt, monkeypatch):
    for key, value in SMTP_ENV.items():
        monkeypatch.setenv(key, value)
    failing = Recorder(error=DeliveryError('Could not send the email.'))
    assert _private(private_root, '--send', sender=failing) == 3


def test_customer_names_never_appear_in_outputs(
        private_root, inbox_file, with_salt, workbook_path, capsys):
    names = set(pd.read_excel(workbook_path, sheet_name='Sales')['Customer'].dropna())
    assert _private(private_root) == 0
    text = capsys.readouterr().out
    for path in (private_root / 'output').iterdir():
        text += path.read_text(encoding='utf-8')
    assert not any(name in text for name in names)


def test_excel_lock_files_in_the_inbox_are_ignored(
        private_root, inbox_file, with_salt):
    lock = private_root / 'inbox' / '~$Boutique_Sales.xlsx'
    lock.write_bytes(b'lock')  # newest file, as Excel leaves while open
    assert _private(private_root) == 0
    assert lock.exists()
    assert not inbox_file.exists()


def test_same_day_rerun_does_not_overwrite_or_crash_on_archive(
        private_root, inbox_file, with_salt, workbook_path):
    assert _private(private_root) == 0
    shutil.copy(workbook_path, inbox_file)
    assert _private(private_root) == 0
    assert len(list((private_root / 'archive').glob('*.xlsx'))) == 2


def test_custom_file_with_synthetic_profile_needs_its_own_output_folder(
        workbook_path, capsys):
    assert _run('--profile', 'synthetic', '--file', str(workbook_path)) == 2
    assert '--demo-output' in capsys.readouterr().out


def test_dashboard_is_only_written_when_asked(workbook_path, tmp_path):
    out = tmp_path / 'demo'
    assert _run('--profile', 'synthetic', '--file', str(workbook_path),
                '--demo-output', str(out)) == 0
    assert not list(out.glob('dashboard_*.html'))


def test_dashboard_flag_writes_a_dashboard_next_to_the_report(
        workbook_path, tmp_path):
    out = tmp_path / 'demo'
    assert _run('--profile', 'synthetic', '--file', str(workbook_path),
                '--demo-output', str(out), '--dashboard') == 0
    page = next(out.glob('dashboard_*.html')).read_text(encoding='utf-8')
    assert 'Boutique dashboard' in page
    assert list(out.glob('weekly_*.html'))


def test_private_dashboard_goes_to_the_private_output_folder(
        private_root, inbox_file, with_salt):
    assert _private(private_root, '--dashboard') == 0
    assert list((private_root / 'output').glob('dashboard_*.html'))


def _send_with_dashboard(private_root, monkeypatch):
    for key, value in SMTP_ENV.items():
        monkeypatch.setenv(key, value)
    sender = Recorder()
    assert _private(private_root, '--send', sender=sender) == 0
    return list(sender.sent[0].iter_attachments())


def test_send_attaches_the_dashboard(
        private_root, inbox_file, with_salt, monkeypatch):
    files = _send_with_dashboard(private_root, monkeypatch)
    assert len(files) == 1
    assert files[0].get_filename().startswith('dashboard_')
    assert files[0].get_filename().endswith('.html')
    assert 'Boutique dashboard' in files[0].get_content()


def test_the_attached_dashboard_has_no_customer_names(
        private_root, inbox_file, with_salt, monkeypatch, workbook_path):
    names = set(pd.read_excel(workbook_path, sheet_name='Sales')['Customer'].dropna())
    attached = _send_with_dashboard(private_root, monkeypatch)[0].get_content()
    assert not any(name in attached for name in names)


def test_private_run_logs_recommendations_next_to_the_private_data(
        private_root, inbox_file, with_salt):
    assert _private(private_root) == 0
    with open_log(private_root / 'decisions.sqlite') as conn:
        kinds = {r.kind for r in list_recommendations(conn)}
    assert kinds == {'forecast', 'reorder'}


def test_untrusted_data_writes_nothing_to_the_decision_log(
        private_root, with_salt, make_workbook):
    shutil.copy(make_workbook(problems={'duplicate_id'}),
                private_root / 'inbox' / 'bad.xlsx')
    assert _private(private_root) == 1
    assert not (private_root / 'decisions.sqlite').exists()


def test_synthetic_demo_does_not_write_a_decision_log(workbook_path, tmp_path):
    code = _run('--profile', 'synthetic', '--file', str(workbook_path),
                '--demo-output', str(tmp_path / 'demo'))
    assert code == 0
    assert not list(tmp_path.rglob('*.sqlite'))


def test_decision_log_flag_chooses_the_location(workbook_path, tmp_path):
    log = tmp_path / 'elsewhere' / 'log.sqlite'
    code = _run('--profile', 'synthetic', '--file', str(workbook_path),
                '--demo-output', str(tmp_path / 'demo'), '--decision-log', str(log))
    assert code == 0
    with open_log(log) as conn:
        assert list_recommendations(conn)


def test_a_broken_decision_log_does_not_stop_the_report(
        private_root, inbox_file, with_salt, capsys):
    (private_root / 'decisions.sqlite').write_text('not a database', encoding='utf-8')
    assert _private(private_root) == 0
    assert 'decision log' in capsys.readouterr().out.lower()
    assert list((private_root / 'output').glob('weekly_*.html'))


def test_private_run_scores_recommendations_whose_window_has_passed(
        private_root, inbox_file, with_salt, capsys):
    log = private_root / 'decisions.sqlite'
    with open_log(log) as conn:
        add_recommendations(conn, [Recommendation(
            as_of=date(2026, 8, 1), product='Silk Saree', kind='forecast',
            model='velocity-v1', horizon_weeks=2, suggested_qty=0,
            expected_demand=3.0, confidence='good')])
    assert _private(private_root) == 0
    with open_log(log) as conn:
        assert len(list_outcomes(conn)) == 1
    assert '1 outcome' in capsys.readouterr().out


# --- repeat-customer reminders (names stay in one local file) -----------------

def _reminder_file(private_root):
    return private_root / 'output' / f'reminders_{TODAY}.txt'


def _expected_reminders(workbook_path):
    return build_reminders(load_customer_rows(workbook_path), TODAY)


def test_reminders_are_off_unless_asked_for(private_root, inbox_file, with_salt):
    assert _private(private_root) == 0
    assert not list((private_root / 'output').glob('reminders_*'))
    assert 'Customers to nudge' not in next(
        (private_root / 'output').glob('weekly_*.txt')).read_text(encoding='utf-8')


def test_the_reminders_flag_writes_a_local_list_with_a_privacy_note(
        private_root, inbox_file, with_salt, workbook_path):
    expected = _expected_reminders(workbook_path)
    assert expected.due_count > 0  # the synthetic data has someone due
    assert _private(private_root, '--reminders') == 0
    text = _reminder_file(private_root).read_text(encoding='utf-8')
    assert PRIVACY_NOTE in text
    assert all(line.name in text for line in expected.due)
    assert not inbox_file.exists()  # still archived afterwards


def test_names_stay_out_of_everything_except_the_reminders_file(
        private_root, inbox_file, with_salt, workbook_path, monkeypatch, capsys):
    names = set(pd.read_excel(workbook_path, sheet_name='Sales')['Customer'].dropna())
    for key, value in SMTP_ENV.items():
        monkeypatch.setenv(key, value)
    sender = Recorder()
    code = _private(private_root, '--reminders', '--dashboard', '--send', sender=sender)
    assert code == 0
    seen = capsys.readouterr().out
    for path in (private_root / 'output').iterdir():
        if not path.name.startswith('reminders_'):
            seen += path.read_text(encoding='utf-8')
    for part in sender.sent[0].walk():
        if part.get_content_maintype() == 'text':
            seen += part.get_content()
    seen += (private_root / 'decisions.sqlite').read_bytes().decode('latin-1')
    assert not any(name in seen for name in names)
    assert 'Customers to nudge' in seen  # the email carries the count


def test_the_report_count_matches_the_list(
        private_root, inbox_file, with_salt, workbook_path):
    expected = _expected_reminders(workbook_path)
    assert _private(private_root, '--reminders') == 0
    report = next((private_root / 'output').glob('weekly_*.txt')).read_text(
        encoding='utf-8')
    assert f'{expected.due_count} repeat customer' in report


def _expected_festival_lists(workbook_path):
    sales = load_workbook_data(workbook_path, salt=SALT).sales
    return build_festival_reminders(
        load_customer_rows(workbook_path), load_calendar(DEFAULT_CALENDAR_PATH),
        festival_windows(sales), pd.Timestamp(TODAY))


def test_the_local_file_lists_last_years_festival_buyers(
        private_root, inbox_file, with_salt, workbook_path):
    expected = _expected_festival_lists(workbook_path)
    assert expected  # Durga Puja is days away and the synthetic data has history
    assert _private(private_root, '--reminders') == 0
    text = _reminder_file(private_root).read_text(encoding='utf-8')
    assert expected[0].festival in text and 'last year' in text
    assert expected[0].regulars[0].name in text


def test_the_report_gets_a_festival_count_but_no_names(
        private_root, inbox_file, with_salt, workbook_path):
    expected = _expected_festival_lists(workbook_path)
    assert _private(private_root, '--reminders') == 0
    report = next((private_root / 'output').glob('weekly_*.txt')).read_text(
        encoding='utf-8')
    assert (f'{expected[0].count} customer' in report
            and f'bought at {expected[0].festival} last year' in report)
    assert expected[0].regulars[0].name not in report


def test_reminders_are_refused_for_the_synthetic_profile(
        workbook_path, tmp_path, capsys):
    code = _run('--profile', 'synthetic', '--file', str(workbook_path),
                '--demo-output', str(tmp_path / 'demo'), '--reminders')
    assert code == 2
    out = capsys.readouterr().out
    assert '--reminders' in out and 'private' in out
    assert not list(tmp_path.rglob('reminders_*'))


def test_untrusted_data_writes_no_reminders(
        private_root, with_salt, make_workbook):
    shutil.copy(make_workbook(problems={'duplicate_id'}),
                private_root / 'inbox' / 'bad.xlsx')
    assert _private(private_root, '--reminders') == 1
    assert not list((private_root / 'output').glob('reminders_*'))
